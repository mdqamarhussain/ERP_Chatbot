import streamlit as st
import requests
from io import BytesIO
from PIL import Image
from fpdf import FPDF
from docx import Document
import pymongo
import tempfile

# MongoDB connection setup
mongo_client = pymongo.MongoClient("mongodb://localhost:27017/")
db = mongo_client["erp_db"]
users_collection = db["users"]
documents_collection = db["documents"]  # Collection to store documents

# Streamlit application
st.title("ERP Chatbot - Document Processing")

# Function to create a new user
def create_user(username, password, role, name):
    # Check if the username already exists
    if users_collection.find_one({"username": username}):
        st.error("Username already exists! Please choose a different username.")
        return False

    # Store the user information
    user_data = {
        "username": username,
        "password": password,  # Store the plain text password (consider hashing for production)
        "role": role,
        "name": name
    }
    users_collection.insert_one(user_data)
    st.success(f"User {username} created successfully!")
    return True

# Function to logout
def logout():
    st.session_state.logged_in = False
    st.session_state.username = ""
    st.success("You have been logged out.")

# Function to save uploaded document to MongoDB
def save_document(file, username):
    document_data = {
        "username": username,
        "file": file.read(),  # Read the file as binary
        "file_name": file.name
    }
    documents_collection.insert_one(document_data)
    st.success("Document saved to database!")

# Login section
st.subheader("Login")
username = st.text_input("Username")
password = st.text_input("Password", type="password")

if st.button("Login"):
    # Check for user in MongoDB
    user = users_collection.find_one({"username": username})
    if user:
        if user['password'] == password:  # Directly compare the password for testing
            st.session_state.logged_in = True
            st.session_state.username = user['name']
            st.success(f"Welcome, {st.session_state.username}!")
        else:
            st.error("Invalid username or password")
    else:
        st.error("Invalid username or password")

# Check if the user is logged in
if 'logged_in' not in st.session_state or not st.session_state.logged_in:
    # Sign-up section
    st.subheader("Sign Up")
    new_username = st.text_input("New Username")
    new_password = st.text_input("New Password", type="password")
    role = st.selectbox("Role", ["user", "admin"])
    name = st.text_input("Name")

    if st.button("Sign Up"):
        if new_username and new_password and role and name:
            create_user(new_username, new_password, role, name)
        else:
            st.error("Please fill in all fields.")

    st.stop()  # Stop execution if the user is not logged in

# Logout button
if st.button("Logout"):
    logout()

# File upload section
uploaded_file = st.file_uploader("Upload an image or document", type=["jpg", "jpeg", "png", "pdf", "docx"])

# Initialize session states
if 'extracted_text' not in st.session_state:
    st.session_state.extracted_text = ""
if 'entities' not in st.session_state:
    st.session_state.entities = []
if 'summary' not in st.session_state:
    st.session_state.summary = ""

if uploaded_file is not None:
    # Show the uploaded file
    st.image(uploaded_file, caption='Uploaded Image', use_column_width=True)

    # Save document button
    if st.button("Save Document to Database"):
        save_document(uploaded_file, st.session_state.username)

    # Extract Text button
    if st.button("Extract Text"):
        with st.spinner("Extracting text..."):
            files = {'file': uploaded_file}
            response = requests.post("http://127.0.0.1:5000/extract-text", files=files)
            if response.status_code == 200:
                st.session_state.extracted_text = response.json().get("extracted_text", "")
                st.success("Text extracted successfully!")

    # Display the extracted text
    if st.session_state.extracted_text:
        st.text_area("Extracted Text", value=st.session_state.extracted_text, height=300, key="extracted_text_area")

        # Summarize button
        if st.button("Summarize Text"):
            with st.spinner("Summarizing text..."):
                response = requests.post("http://127.0.0.1:5000/summarize-text", json={"text": st.session_state.extracted_text})
                if response.status_code == 200:
                    st.session_state.summary = response.json().get("summary", "")
                    st.success("Text summarized successfully!")
                else:
                    st.error("Error summarizing text!")

        # Display the summary
        if st.session_state.summary:
            st.text_area("Summary", value=st.session_state.summary, height=150, key="summary_area")

        # Classify Document button for text classification
        if st.button("Classify Document"):
            with st.spinner("Classifying text..."):
                response = requests.post("http://127.0.0.1:5000/classify-text", json={"text": st.session_state.extracted_text})
                if response.status_code == 200:
                    classification = response.json().get("classification", "Unknown")
                    st.success(f"Document classified as: {classification}")
                else:
                    st.error("Error classifying document!")

        # Dropdown menu for conversion options
        conversion_option = st.selectbox("Choose a conversion option:", ["PDF", "Word"])

        if conversion_option == "PDF":
            pdf_option = st.selectbox("Choose PDF conversion type:", ["Extracted Text to PDF", "Image to PDF"])

            if pdf_option == "Extracted Text to PDF":
                if st.button("Download as PDF"):
                    # Create a temporary file to save the PDF
                    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as temp_pdf:
                        pdf = FPDF()
                        pdf.add_page()
                        pdf.set_font("Arial", size=12)
                        for line in st.session_state.extracted_text.split('\n'):
                            pdf.cell(200, 10, txt=line, ln=True)

                        # Save the PDF to the temporary file
                        pdf.output(temp_pdf.name)

                        # Prepare for download
                        with open(temp_pdf.name, 'rb') as f:
                            pdf_output = f.read()

                        # Direct download
                        st.download_button(
                            label="Download PDF",
                            data=pdf_output,
                            file_name="extracted_text.pdf",
                            mime="application/pdf",
                            key="pdf_download"
                        )

            elif pdf_option == "Image to PDF":
                if st.button("Download as PDF"):
                    img = Image.open(uploaded_file)
                    pdf_output = BytesIO()
                    img.save(pdf_output, "PDF", resolution=100.0)
                    pdf_output.seek(0)

                    # Direct download
                    st.download_button(
                        label="Download PDF",
                        data=pdf_output,
                        file_name="image.pdf",
                        mime="application/pdf",
                        key="image_pdf_download"
                    )

        elif conversion_option == "Word":
            if st.button("Download as Word"):
                # Create a temporary file to save the Word document
                with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as temp_doc:
                    doc = Document()
                    doc.add_paragraph(st.session_state.extracted_text)

                    # Save the Word document to the temporary file
                    doc.save(temp_doc.name)

                    # Prepare for download
                    with open(temp_doc.name, 'rb') as f:
                        word_output = f.read()

                    # Direct download
                    st.download_button(
                        label="Download Word",
                        data=word_output,
                        file_name="extracted_text.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        key="word_download"
                    )